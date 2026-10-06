#!/usr/bin/env python
"""BENCH-001: Pandas vs Polars (+ DuckDB) for ingestion, profiling, aggregation.

Feeds ADR-0004 (Parquet + locked-down DuckDB) and ADR-0005 (DataFrame library).
See docs/research/RES-014.md ("Open items for BENCH-001") and docs/trd.md
§11.1 (profile fields) / §15.2 (operations catalog).

Run (exact versions RES-014 recommends):

    uv run --python 3.14 --with polars==1.44.2 --with pandas==3.0.6 \
        --with duckdb==1.5.5 --with pyarrow==25.0.1 --with numpy==2.5.3 \
        --with psutil --with openpyxl==3.1.5 --with python-calamine==0.8.2 \
        python scripts/bench/bench_001.py

All CSV/Parquet/xlsx files are written under a temp dir outside the repo
(tempfile.gettempdir()) and never committed. Row-count tiers can be
overridden for a quick smoke test with BENCH_TIERS="1000,5000".

Every number in the printed report is measured in this process. Nothing is
invented; a failed measurement is recorded as such (see EXCEL section and
DuckDB Mode B "works" flag).
"""
from __future__ import annotations

import gc
import json
import os
import sys
import tempfile
import threading
import time
import traceback
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import polars as pl
import pyarrow as pa
import pyarrow.dataset as pa_dataset
import pyarrow.parquet as pq
import duckdb

try:
    import psutil

    HAVE_PSUTIL = True
except ImportError:
    import tracemalloc

    HAVE_PSUTIL = False

SEED = 20260928
N_WARMUP = 1
N_RUNS = 3
CATEGORIES = ["Electronics", "Grocery", "Apparel", "Home", "Toys", "Sports", "Beauty", "Books"]
REGIONS = ["North", "South", "East", "West", "Central"]

TIERS = [int(x) for x in os.environ.get("BENCH_TIERS", "10000,100000,1000000").split(",")]
XLSX_ROWS = 100_000

DATA_DIR = Path(tempfile.mkdtemp(prefix="bench001_data_"))
print(f"# data dir (outside repo, not committed): {DATA_DIR}", file=sys.stderr)


# --------------------------------------------------------------------------- #
# Synthetic retail-like dataset generator (fixed seed -> reproducible)
# --------------------------------------------------------------------------- #
def gen_dataframe(n: int, seed: int = SEED) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    order_id = np.arange(1, n + 1)
    customer_id = rng.integers(1, max(2, n // 5), size=n)
    start = date(2023, 1, 1)
    day_offsets = rng.integers(0, 730, size=n)
    order_date = np.array([(start + timedelta(days=int(d))).isoformat() for d in day_offsets])
    category = rng.choice(CATEGORIES, size=n).astype(object)
    region = rng.choice(REGIONS, size=n)
    quantity = rng.integers(1, 20, size=n)
    unit_price = np.round(rng.uniform(1, 500, size=n), 2)
    noise = rng.normal(1.0, 0.02, size=n)
    amount = np.round(quantity * unit_price * noise, 2).astype(object)
    discount_pct = np.round(rng.uniform(0, 0.3, size=n), 4).astype(object)
    is_returned = rng.random(n) < 0.07

    # inject nulls (retail data is messy)
    amount[rng.random(n) < 0.03] = None
    discount_pct[rng.random(n) < 0.05] = None
    category[rng.random(n) < 0.01] = None

    return pd.DataFrame(
        {
            "order_id": order_id,
            "customer_id": customer_id,
            "order_date": order_date,
            "category": category,
            "region": region,
            "quantity": quantity,
            "unit_price": unit_price,
            "amount": amount,
            "discount_pct": discount_pct,
            "is_returned": is_returned,
        }
    )


# --------------------------------------------------------------------------- #
# Timing + peak-RSS measurement
# --------------------------------------------------------------------------- #
def _rss() -> int:
    return psutil.Process(os.getpid()).memory_info().rss


class _MemSampler:
    """ponytail: coarse whole-process RSS polling thread, not a per-allocation
    profiler. Good enough to compare libraries' order-of-magnitude memory use;
    upgrade to a proper allocator hook if finer precision is ever needed.
    """

    def __init__(self, interval: float = 0.005):
        self.interval = interval
        self.peak = 0
        self._stop = threading.Event()
        self._thread = None

    def __enter__(self):
        if HAVE_PSUTIL:
            self.peak = _rss()
            self._stop.clear()

            def run():
                while not self._stop.is_set():
                    self.peak = max(self.peak, _rss())
                    time.sleep(self.interval)

            self._thread = threading.Thread(target=run, daemon=True)
            self._thread.start()
        else:
            tracemalloc.start()
        return self

    def __exit__(self, *exc):
        if HAVE_PSUTIL:
            self._stop.set()
            self._thread.join()
        else:
            _, peak = tracemalloc.get_traced_memory()
            self.peak = peak
            tracemalloc.stop()


def measure(fn, label: str, warmup: int = N_WARMUP, runs: int = N_RUNS):
    """Median wall time (perf_counter) and median peak RSS delta over `runs`
    timed calls, after `warmup` untimed calls. Returns (stats_dict, last_result).
    """
    result = None
    for _ in range(warmup):
        result = fn()
    del result
    gc.collect()

    times = []
    peaks = []
    result = None
    for _ in range(runs):
        del result
        gc.collect()
        base = _rss() if HAVE_PSUTIL else 0
        with _MemSampler() as sampler:
            t0 = time.perf_counter()
            result = fn()
            t1 = time.perf_counter()
        times.append(t1 - t0)
        peaks.append(max(0, sampler.peak - base) if HAVE_PSUTIL else sampler.peak)

    times_sorted = sorted(times)
    peaks_sorted = sorted(peaks)
    mid = len(times_sorted) // 2
    med_time = times_sorted[mid] if len(times_sorted) % 2 else (times_sorted[mid - 1] + times_sorted[mid]) / 2
    med_peak = peaks_sorted[len(peaks_sorted) // 2]
    stats = {
        "label": label,
        "median_s": round(med_time, 6),
        "runs_s": [round(t, 6) for t in times],
        "median_peak_mb": round(med_peak / (1024 * 1024), 2),
        "peaks_mb": [round(p / (1024 * 1024), 2) for p in peaks],
        "mem_method": "psutil_rss_delta" if HAVE_PSUTIL else "tracemalloc",
    }
    return stats, result


# --------------------------------------------------------------------------- #
# Operations under test
# --------------------------------------------------------------------------- #
def read_csv_pandas(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, parse_dates=["order_date"])


def read_csv_polars(path: Path) -> pl.DataFrame:
    return pl.read_csv(path, try_parse_dates=True)


def write_parquet_from_pandas(df: pd.DataFrame, out: Path) -> None:
    table = pa.Table.from_pandas(df, preserve_index=False)
    pq.write_table(table, out, compression="zstd")


def write_parquet_from_polars(df: pl.DataFrame, out: Path) -> None:
    table = df.to_arrow()
    pq.write_table(table, out, compression="zstd")


def _f(x):
    try:
        return None if pd.isna(x) else float(x)
    except TypeError:
        return None


def profile_pandas(df: pd.DataFrame) -> dict:
    out = {}
    for col in df.columns:
        s = df[col]
        entry = {"null_ratio": float(s.isna().mean()), "distinct": int(s.nunique(dropna=True))}
        if pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s):
            entry.update(
                {
                    "min": _f(s.min()),
                    "max": _f(s.max()),
                    "mean": _f(s.mean()),
                    "std": _f(s.std()),
                    "q25": _f(s.quantile(0.25)),
                    "q50": _f(s.quantile(0.5)),
                    "q75": _f(s.quantile(0.75)),
                }
            )
        elif not pd.api.types.is_datetime64_any_dtype(s):
            entry["top10"] = {str(k): int(v) for k, v in s.value_counts(dropna=True).head(10).items()}
        out[col] = entry
    return out


NUMERIC_POLARS_DTYPES = {pl.Int64, pl.Int32, pl.Float64, pl.Float32}


def profile_polars(df: pl.DataFrame) -> dict:
    n = df.height
    numeric_cols = [c for c, dt in zip(df.columns, df.dtypes) if dt in NUMERIC_POLARS_DTYPES]
    other_cols = [
        c for c in df.columns if c not in numeric_cols and df.schema[c] not in (pl.Datetime, pl.Date)
    ]
    exprs = []
    for c in df.columns:
        exprs.append(pl.col(c).null_count().alias(f"{c}__nulls"))
        exprs.append(pl.col(c).n_unique().alias(f"{c}__distinct"))
    for c in numeric_cols:
        exprs += [
            pl.col(c).min().alias(f"{c}__min"),
            pl.col(c).max().alias(f"{c}__max"),
            pl.col(c).mean().alias(f"{c}__mean"),
            pl.col(c).std().alias(f"{c}__std"),
            pl.col(c).quantile(0.25).alias(f"{c}__q25"),
            pl.col(c).quantile(0.5).alias(f"{c}__q50"),
            pl.col(c).quantile(0.75).alias(f"{c}__q75"),
        ]
    row = df.select(exprs).row(0, named=True)
    out = {}
    for c in df.columns:
        entry = {"null_ratio": row[f"{c}__nulls"] / n, "distinct": row[f"{c}__distinct"]}
        if c in numeric_cols:
            entry.update({k: row.get(f"{c}__{k}") for k in ["min", "max", "mean", "std", "q25", "q50", "q75"]})
        out[c] = entry
    for c in other_cols:
        top10 = df[c].value_counts().sort("count", descending=True).head(10)
        out[c]["top10"] = dict(zip([str(v) for v in top10[c].to_list()], top10["count"].to_list()))
    return out


def groupby_pandas(df: pd.DataFrame):
    d = df[["category", "order_date", "amount"]].copy()
    d["month"] = d["order_date"].dt.to_period("M").astype(str)
    return d.groupby(["category", "month"], dropna=False).agg(total_amount=("amount", "sum"), n=("amount", "count"))


def groupby_polars(df: pl.DataFrame):
    return (
        df.select(["category", "order_date", "amount"])
        .with_columns(pl.col("order_date").dt.strftime("%Y-%m").alias("month"))
        .group_by(["category", "month"])
        .agg(pl.col("amount").sum().alias("total_amount"), pl.col("amount").count().alias("n"))
    )


# --------------------------------------------------------------------------- #
# DuckDB lockdown recipes (RES-014 R3 / ADR-0004)
# --------------------------------------------------------------------------- #
LOCK_CONFIG = {
    "autoload_known_extensions": "false",
    "autoinstall_known_extensions": "false",
    "allow_community_extensions": "false",
}


def duckdb_mode_a_query(parquet_path: Path, filtered: bool):
    config = dict(LOCK_CONFIG)
    config["enable_external_access"] = "false"
    con = duckdb.connect(":memory:", config=config)
    try:
        ds = pa_dataset.dataset(str(parquet_path), format="parquet")
        con.register("sales", ds)
        con.execute("SET lock_configuration=true")
        sql = (
            "SELECT category, date_trunc('month', order_date) AS month, "
            "SUM(amount) AS total, COUNT(*) AS n FROM sales"
        )
        if filtered:
            sql += " WHERE category = 'Electronics' AND amount > 100"
        sql += " GROUP BY 1, 2"
        return con.execute(sql).fetchall()
    finally:
        con.close()


def duckdb_mode_b_query(parquet_path: Path, filtered: bool):
    config = dict(LOCK_CONFIG)
    con = duckdb.connect(":memory:", config=config)
    try:
        p = parquet_path.resolve().as_posix()
        con.execute("SET allowed_paths=[?]", [p])
        con.execute("SET enable_external_access=false")
        con.execute(f"CREATE VIEW sales AS SELECT * FROM read_parquet('{p}')")
        con.execute("SET lock_configuration=true")
        sql = (
            "SELECT category, date_trunc('month', order_date) AS month, "
            "SUM(amount) AS total, COUNT(*) AS n FROM sales"
        )
        if filtered:
            sql += " WHERE category = 'Electronics' AND amount > 100"
        sql += " GROUP BY 1, 2"
        return con.execute(sql).fetchall()
    finally:
        con.close()


def duckdb_mode_a_profile(parquet_path: Path) -> dict:
    config = dict(LOCK_CONFIG)
    config["enable_external_access"] = "false"
    con = duckdb.connect(":memory:", config=config)
    try:
        ds = pa_dataset.dataset(str(parquet_path), format="parquet")
        con.register("sales", ds)
        con.execute("SET lock_configuration=true")
        summary = con.execute("SUMMARIZE SELECT * FROM sales").fetchall()
        cols = [d[0] for d in con.description]
        out = {"summarize_columns": cols, "summarize_rows": [list(r) for r in summary]}
        for c in ("category", "region", "is_returned"):
            top = con.execute(
                f"SELECT {c}, COUNT(*) AS n FROM sales GROUP BY {c} ORDER BY n DESC LIMIT 10"
            ).fetchall()
            out[f"top10_{c}"] = top
        return out
    finally:
        con.close()


def duckdb_mode_a_smoke_check(parquet_path: Path) -> str:
    """Confirm Mode A actually works and has no filesystem/network reach."""
    config = dict(LOCK_CONFIG)
    config["enable_external_access"] = "false"
    con = duckdb.connect(":memory:", config=config)
    try:
        ds = pa_dataset.dataset(str(parquet_path), format="parquet")
        con.register("sales", ds)
        con.execute("SET lock_configuration=true")
        n = con.execute("SELECT COUNT(*) FROM sales").fetchone()[0]
        try:
            con.execute(f"SELECT * FROM read_parquet('{parquet_path.as_posix()}')")
            blocked = False
        except Exception:
            blocked = True
        return f"OK: registered table has {n} rows; direct read_parquet blocked={blocked}"
    except Exception as e:
        return f"FAILED: {e!r}"
    finally:
        con.close()


# --------------------------------------------------------------------------- #
# Excel: openpyxl (read_only, values_only) vs python-calamine
# --------------------------------------------------------------------------- #
def make_xlsx(n: int, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("orders")
    ws.append(["order_id", "customer_id", "order_date", "category", "region", "quantity", "unit_price", "amount"])
    df = gen_dataframe(n)
    cols = ["order_id", "customer_id", "order_date", "category", "region", "quantity", "unit_price", "amount"]
    for row in df[cols].itertuples(index=False):
        ws.append(list(row))
    wb.save(path)


def read_xlsx_openpyxl(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def read_xlsx_calamine(path: Path):
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(str(path))
    sheet = wb.get_sheet_by_index(0)
    return sheet.to_python()


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    report = {
        "seed": SEED,
        "tiers": TIERS,
        "have_psutil": HAVE_PSUTIL,
        "versions": {
            "python": sys.version,
            "polars": pl.__version__,
            "pandas": pd.__version__,
            "duckdb": duckdb.__version__,
            "pyarrow": pa.__version__,
            "numpy": np.__version__,
            "psutil": psutil.__version__ if HAVE_PSUTIL else None,
        },
        "tier_results": {},
        "excel": {},
    }

    for name, ver in report["versions"].items():
        print(f"# version {name}: {ver}", file=sys.stderr)

    for n in TIERS:
        print(f"\n### tier {n} rows ###", file=sys.stderr)
        tier = {}
        csv_path = DATA_DIR / f"orders_{n}.csv"
        gen_dataframe(n).to_csv(csv_path, index=False)
        tier["csv_bytes"] = csv_path.stat().st_size

        print("  csv parse ...", file=sys.stderr)
        tier["read_csv_pandas"], df_pd = measure(lambda: read_csv_pandas(csv_path), "pandas.read_csv")
        tier["read_csv_polars"], df_pl = measure(lambda: read_csv_polars(csv_path), "polars.read_csv")

        pq_pd_path = DATA_DIR / f"pd_{n}.parquet"
        pq_pl_path = DATA_DIR / f"pl_{n}.parquet"
        print("  parquet write ...", file=sys.stderr)
        tier["write_parquet_pandas_df"], _ = measure(
            lambda: write_parquet_from_pandas(df_pd, pq_pd_path), "pandas_df->pyarrow.parquet"
        )
        tier["write_parquet_polars_df"], _ = measure(
            lambda: write_parquet_from_polars(df_pl, pq_pl_path), "polars_df->pyarrow.parquet"
        )
        tier["parquet_bytes_from_pandas_df"] = pq_pd_path.stat().st_size
        tier["parquet_bytes_from_polars_df"] = pq_pl_path.stat().st_size

        print("  profiling ...", file=sys.stderr)
        tier["profile_pandas"], _ = measure(lambda: profile_pandas(df_pd), "pandas profile")
        tier["profile_polars"], _ = measure(lambda: profile_polars(df_pl), "polars profile")
        tier["profile_duckdb_mode_a"], _ = measure(lambda: duckdb_mode_a_profile(pq_pd_path), "duckdb SUMMARIZE (Mode A)")

        print("  groupby ...", file=sys.stderr)
        tier["groupby_pandas"], _ = measure(lambda: groupby_pandas(df_pd), "pandas groupby")
        tier["groupby_polars"], _ = measure(lambda: groupby_polars(df_pl), "polars group_by")

        print("  duckdb query modes ...", file=sys.stderr)
        tier["duckdb_mode_a_smoke"] = duckdb_mode_a_smoke_check(pq_pd_path)
        tier["duckdb_mode_a_agg"], _ = measure(lambda: duckdb_mode_a_query(pq_pd_path, False), "DuckDB Mode A group-by")
        tier["duckdb_mode_a_filtered"], _ = measure(
            lambda: duckdb_mode_a_query(pq_pd_path, True), "DuckDB Mode A filtered"
        )
        try:
            tier["duckdb_mode_b_agg"], _ = measure(
                lambda: duckdb_mode_b_query(pq_pd_path, False), "DuckDB Mode B group-by"
            )
            tier["duckdb_mode_b_filtered"], _ = measure(
                lambda: duckdb_mode_b_query(pq_pd_path, True), "DuckDB Mode B filtered"
            )
            tier["duckdb_mode_b_status"] = "OK"
        except Exception as e:
            tier["duckdb_mode_b_status"] = f"FAILED: {e!r}\n{traceback.format_exc()}"

        report["tier_results"][str(n)] = tier
        del df_pd, df_pl
        gc.collect()

    print("\n### excel ###", file=sys.stderr)
    xlsx_path = DATA_DIR / f"orders_{XLSX_ROWS}.xlsx"
    make_xlsx(XLSX_ROWS, xlsx_path)
    report["excel"]["rows"] = XLSX_ROWS
    report["excel"]["xlsx_bytes"] = xlsx_path.stat().st_size
    try:
        report["excel"]["openpyxl"], res = measure(
            lambda: read_xlsx_openpyxl(xlsx_path), "openpyxl read_only+values_only", runs=N_RUNS
        )
        report["excel"]["openpyxl_status"] = f"OK ({len(res)} rows read)"
    except Exception as e:
        report["excel"]["openpyxl_status"] = f"FAILED: {e!r}\n{traceback.format_exc()}"

    try:
        report["excel"]["calamine"], res = measure(
            lambda: read_xlsx_calamine(xlsx_path), "python-calamine", runs=N_RUNS
        )
        report["excel"]["calamine_status"] = f"OK ({len(res)} rows read)"
    except Exception as e:
        report["excel"]["calamine_status"] = f"FAILED: {e!r}\n{traceback.format_exc()}"

    print("\n===JSON_RESULTS_START===")
    print(json.dumps(report, indent=2, default=str))
    print("===JSON_RESULTS_END===")


if __name__ == "__main__":
    main()
